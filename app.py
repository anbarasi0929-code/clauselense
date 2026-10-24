from flask import Flask, request, jsonify

from flask_cors import CORS

from db import db

from pathlib import Path

from docx import Document

import pymupdf

import io

import re
from datetime import datetime, timezone

app = Flask(__name__)

CORS(app)

app.config["MAX_CONTENT_LENGTH"] = 10 * 1024 * 1024

ALLOWED_EXTENSIONS = {".pdf", ".docx"}

# Built-in CLM-style screening rules.

# These identify possible issues, not definitive legal conclusions.

CLM_RULES = [

{

    "category": "Termination",

    "severity": "High",

    "patterns": [

        r"\btermination for convenience\b",

        r"\btermination for cause\b",

        r"\bterminate(?:d|s|ing)?\s+(?:this|the)\s+agreement\b",

        r"\bend this agreement\b",

        r"\bnotice of termination\b",

        r"\btermination notice period\b"

    ],

    "reason": "Review the grounds for termination, required notice, and each party's exit rights.",

    "recommendation": "Review termination rights, notice periods, and whether both parties have fair termination options."

},

    {

        "category": "Payment",

        "severity": "Medium",

        "patterns": [

            r"\bpayment\s+(?:shall|must|will|is due|due)\b",

            r"\bpayment terms\b",

            r"\bpayment deadline\b",

            r"\binvoices?\s+(?:shall|must|will|are|is)\b",

            r"\b(?:late fee|late payment fee)\b",

            r"\binterest on overdue payments?\b",

            r"\bsuspend(?:ed|sion)?\s+(?:the\s+)?services?\s+for\s+nonpayment\b"

        ],

        "reason": "Review payment deadlines, late fees, invoice disputes, and the consequences of nonpayment.",

"recommendation": "Check payment due dates, late charges, invoice dispute procedures, and what happens if payment is delayed."

},

    {

        "category": "Liability",

        "severity": "High",

        "patterns": [

            r"\blimitation of liability\b",

            r"\blimit(?:ed|ation)?\s+(?:of\s+)?liability\b",

            r"\bliability cap\b",

            r"\bliable for (?:any|all|direct|indirect)\b",

            r"\bconsequential damages\b",

            r"\bindirect damages\b",

            r"\bexclusion of damages\b"

        ],

        "reason": "Review the liability limits, excluded damages, and each party's responsibility for losses.",

"recommendation": "Check the liability cap, excluded types of damages, and whether each party's responsibilities are reasonable."

    },

    {

        "category": "Indemnity",

        "severity": "High",

        "patterns": [

            r"\bindemnif(?:y|ies|ied|ication|ying)\b",

            r"\bindemnify and hold harmless\b",

            r"\bhold harmless\b",

            r"\bduty to defend\b"

        ],

       "reason": "Review who must cover claims, losses, and legal expenses under the agreement.",

"recommendation": "Check which claims are covered, who controls the defense, and whether the indemnity obligations have reasonable limits."

    },

    {

        "category": "Confidentiality",

        "severity": "Medium",

        "patterns": [

            r"\bconfidential information\b",

            r"\bconfidentiality obligations?\b",

            r"\bnon[- ]disclosure\b",

            r"\bproprietary information\b",

            r"\btrade secrets?\b"

        ],

        "reason": "Review what information is protected, permitted disclosures, exceptions, and how long confidentiality obligations last.",

"recommendation": "Check which information is confidential, who may receive it, what exceptions apply, and when the confidentiality obligations end."

},

    {

        "category": "Automatic Renewal",

        "severity": "Medium",

        "patterns": [

            r"\bautomatically renew(?:s|ed|ing)?\b",

            r"\bautomatic renewal\b",

            r"\brenewal term\b",

            r"\brenewal period\b",

            r"\bunless either party gives notice\b",

            r"\brenew(?:s|ed|al)?\s+for successive terms\b"

        ],

        "reason": "Review the renewal period, renewal dates, and deadline for giving cancellation notice.",

"recommendation": "Check when the agreement renews, how long each renewal lasts, and the deadline and method for cancelling before renewal."

},

    {

        "category": "Dispute Resolution",

        "severity": "Medium",

        "patterns": [

            r"\bdispute resolution\b",

            r"\bbinding arbitration\b",

            r"\bsubmit(?:ted)?\s+to arbitration\b",

            r"\bmediation shall\b",

            r"\bgoverning law\b",

            r"\bexclusive jurisdiction\b",

            r"\bexclusive venue\b"

        ],

        "reason": "Review the procedures for resolving disputes, applicable law, and the court or arbitration location.",

        "recommendation": "Check whether disputes require mediation or arbitration, which law applies, and where disputes must be handled."

},

    {

        "category": "Unilateral Changes",

        "severity": "High",

        "patterns": [

            r"\bsole discretion to (?:modify|change|amend)\b",

            r"\bwithout prior notice\b",

            r"\bmodify this agreement without\b",

            r"\bchange these terms without\b",

            r"\breserves? the right to (?:modify|change|amend)\b",

            r"\bunilaterally (?:modify|change|amend)\b"

        ],

        "reason": "Review whether one party can change the agreement without the other party's consent or prior notice.",

"recommendation": "Check whether advance notice is required, whether both parties must agree to changes, and whether you can reject proposed changes."

},

]

def extract_text(file):

    filename = file.filename.lower()

    content = file.read()

    if filename.endswith(".pdf"):

        with pymupdf.open(stream=content, filetype="pdf") as pdf:

            pages = []

            for page_number, page in enumerate(pdf, start=1):

                page_text = page.get_text().strip()

                if page_text:

                    pages.append(

                        f"[PAGE {page_number}]\n{page_text}"

                    )

            return "\n".join(pages)

    if filename.endswith(".docx"):

        document = Document(io.BytesIO(content))

        paragraphs = [p.text for p in document.paragraphs]

        # Include text from tables too.

        for table in document.tables:

            for row in table.rows:

                paragraphs.append(

                    " | ".join(cell.text for cell in row.cells)

                )

        return "\n".join(paragraphs)

    raise ValueError("Only PDF and DOCX files are supported.")

def split_into_clauses(text):

    page_pattern = r"\[PAGE\s+(\d+)\]"

    page_parts = re.split(page_pattern, text)

    clauses = []

    for i in range(1, len(page_parts), 2):

        page_number = page_parts[i]

        page_text = page_parts[i + 1]

        sections = re.split(

            r"\n\s*\n|(?<=[.!?])\s+(?=[A-Z])",

            page_text

        )

        for section in sections:

            section = section.strip()

            if not section:

                continue

            section = re.sub(r"[•●▪◦]", "", section)

            section = re.sub(r"\s+\d+\.\s*$", "", section)

            section = section.strip()

            if not section:

                continue

            if re.match(

                r"^(PAGE\s+\d+|SAMPLE SERVICE AGREEMENT)",

                section,

                re.IGNORECASE

            ):

                continue

            if len(section.split()) < 5:

                continue

            clauses.append(f"[PAGE {page_number}]\n{section}")

    return clauses

def detect_risks(text):

    findings = []

    seen = set()

    for clause in split_into_clauses(text):

        page_match = re.match(r"^\[PAGE\s+(\d+)\]\s*", clause)

        current_page = None

        if page_match:

            current_page = int(page_match.group(1))

            clause = clause[page_match.end():]

        normalized_clause = re.sub(r"\s+", " ", clause).strip()

        if not normalized_clause:

            continue

        for rule in CLM_RULES:

            matched = any(

                re.search(pattern, normalized_clause, re.IGNORECASE)

                for pattern in rule["patterns"]

            )

            if not matched:

                continue

            normalized_key = re.sub(

                r"[\W_]+", " ", normalized_clause.lower()

            ).strip()

            key = (rule["category"], normalized_key)

            if key in seen:

                continue

            findings.append({

                "category": rule["category"],

                "severity": rule["severity"],

                "clause": normalized_clause[:1500],

                "page": current_page,

                "reason": rule.get("reason", "Review this clause carefully."),

                "recommendation": rule.get(

                    "recommendation",

                    "Review this clause against your requirements."

                ),

                "source": "Built-in CLM rule"

            })

            seen.add(key)

    return findings

@app.route("/")

def home():

    return jsonify({

        "message": "ClauseLens backend is running!"

    })

@app.route("/api/analyze", methods=["POST"])

def analyze_contract():

    if "file" not in request.files:

        return jsonify({

            "error": "Please upload a contract."

        }), 400

    file = request.files["file"]

    if not file.filename:

        return jsonify({

            "error": "Please select a file."

        }), 400

    extension = Path(file.filename).suffix.lower()

    if extension not in ALLOWED_EXTENSIONS:

        return jsonify({

            "error": "Only PDF and DOCX files are supported."

        }), 400

    try:

        text = extract_text(file)

        print("PAGE MARKERS:", re.findall(r"\[PAGE\s+\d+\]", text))

        if not text.strip():

            return jsonify({

                "error": "No readable text found. Scanned PDFs may need OCR."

            }), 400

        findings = detect_risks(text)

        print("FINDING PAGES:", [f["page"] for f in findings])

        # Optional user requirements. They are returned for later

        # implementation; the built-in rules run regardless.

        requirements = request.form.get("requirements", "").strip()

        analysis_record = {
            "filename": file.filename,
            "character_count": len(text),
            "requirements": requirements,
            "analysis_type": "Built-in CLM-style rule screening",
            "total_findings": len(findings),
            "findings": findings,
            "created_at": datetime.now(timezone.utc)
        }
        result = db.contract_analyses.insert_one(analysis_record)

        return jsonify({

            "analysis_id": str(result.inserted_id),
            "filename": file.filename,

            "message": "Contract screening completed.",

            "character_count": len(text),

            "requirements_provided": bool(requirements),

            "custom_requirements": requirements,

            "analysis_type": "Built-in CLM-style rule screening",

            "total_findings": len(findings),

            "findings": findings,

            "disclaimer": (

                "Automated preliminary screening only. "

                "Flags are not legal conclusions; review the full contract."

            )

        })

    except Exception:

        app.logger.exception("Contract analysis failed")

        return jsonify({

            "error": "Unable to read or analyze this document."

        }), 500

@app.errorhandler(413)

def file_too_large(error):

    return jsonify({

        "error": "File is too large. Maximum upload size is 10 MB."

    }), 413

@app.route("/db-test", methods=["GET"])

def db_test():

    try:

        result = db.connection_test.insert_one({

            "project": "ClauseLens",

            "message": "Flask connected to MongoDB"

        })

        return jsonify({

            "status": "success",

            "message": "Flask connected to MongoDB Atlas",

            "inserted_id": str(result.inserted_id)

        }), 200

    except Exception:

        app.logger.exception("MongoDB operation failed")

        return jsonify({

            "status": "error",

            "message": "Database operation failed"

        }), 500
@app.route("/api/history", methods=["GET"])
def get_analysis_history():
    try:
        records = list(
            db.contract_analyses
            .find({})
            .sort("created_at", -1)
        )

        reviews = []

        for record in records:
            findings = record.get("findings", [])
            high_count = sum(
                1 for finding in findings
                if finding.get("severity", "").lower() == "high"
            )

            reviews.append({
                "id": str(record["_id"]),
                "filename": record.get("filename", "Unnamed contract"),
                "review_date": (
                    record["created_at"].isoformat()
                    if record.get("created_at") else None
                ),
                "flags": record.get("total_findings", len(findings)),
                "status": (
                    "Needs Attention" if high_count > 0
                    else "Completed"
                ),
                "character_count": record.get("character_count", 0),
                "findings": findings
            })

        return jsonify({
            "total_reviews": len(reviews),
            "completed_reviews": sum(
                1 for review in reviews
                if review["status"] == "Completed"
            ),
            "needs_attention": sum(
                1 for review in reviews
                if review["status"] == "Needs Attention"
            ),
            "reviews": reviews
        }), 200

    except Exception:
        app.logger.exception("Could not load review history")
        return jsonify({"error": "Could not load review history"}), 500
if __name__ == "__main__":

    app.run(debug=True)
